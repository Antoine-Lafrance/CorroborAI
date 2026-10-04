from pathlib import Path

import pytest

from corroborai import pipeline, report
from corroborai.tree import AMBIGU, ANOMALIE, CONFORME, ECART_JUSTIFIE

DATA = Path(__file__).resolve().parents[1] / "data" / "raw"


@pytest.fixture(scope="module")
def res():
    return pipeline.run(DATA, overrides=[])


def test_ai_layer_resolves_every_ambiguous_row(res):
    df = res.df
    amb = df[df.VerdictDeterministe == AMBIGU]
    assert len(amb) == 52
    assert (amb.Couche == "IA").all() and (amb.Verdict != AMBIGU).all()
    assert set(amb.ChampCible) == {"contactEmail", "positionName", "weeklyHoursOverride", "dailyHoursOverride"}
    assert amb.Justification.str.startswith("IA :").all()
    assert amb.Chemin.str.contains(r"\[IA\.").all()


def test_ai_does_not_touch_deterministic_verdicts(res):
    det = res.df[res.df.Couche == "DETERMINISTE"]
    assert (det.Verdict == det.VerdictDeterministe).all()
    assert (det.Confiance == 1.0).all()


def test_hours_rule_is_learned(res):
    m = res.models["weeklyHoursOverride"]
    assert m.reliable and m.rule == "poste.HeuresSemaineContrat" and m.coverage == 1.0


def test_anomalies_are_prioritised(res):
    bad = res.df[res.df.Verdict == ANOMALIE]
    assert len(bad) == 10 and (bad.Priorite > 0).all()
    assert bad.iloc[0].ChampCible == "(affectation)"  # missing assignment ranks first
    assert (res.df[res.df.Verdict.isin([CONFORME, ECART_JUSTIFIE])].Priorite == 0).all()


def test_expert_override_by_pattern_generalises(res):
    row = res.df[(res.df.ChampCible == "contractTypeCode") & (res.df.Verdict == ANOMALIE)].iloc[0]
    o = {"date": "2026-01-01T00:00:00", "champ": row.ChampCible, "signature": row.Signature,
         "portee": "motif", "verdict": ECART_JUSTIFIE, "commentaire": "test"}
    df = pipeline.apply_overrides(res.df, [o])
    hit = df[(df.ChampCible == row.ChampCible) & (df.Signature == row.Signature)]
    assert len(hit) > 1 and (hit.Verdict == ECART_JUSTIFIE).all() and (hit.Couche == "EXPERT").all()


def test_demo_has_the_three_required_cases(res):
    demo = pipeline.demo_cases(res.df)
    assert demo["conforme"]["Verdict"] == CONFORME
    assert demo["ecart_justifie"]["Verdict"] == ECART_JUSTIFIE
    assert demo["anomalie"]["Verdict"] == ANOMALIE


def test_report_sheets(res, tmp_path):
    import openpyxl
    path = report.write(res.df, tmp_path, pipeline.learned_rules(res), [])
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ["Résumé", "À investiguer"]
    assert "Règles apprises (IA)" in wb.sheetnames
    assert wb["À investiguer"].max_row == 11
